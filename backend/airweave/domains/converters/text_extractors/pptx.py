"""Direct text extraction from PPTX files using python-pptx."""

from __future__ import annotations

import asyncio
import os
from collections.abc import Iterator
from itertools import chain
from typing import TYPE_CHECKING, Optional, cast

from airweave.core.logging import logger
from airweave.domains.converters._base import ConversionResult
from airweave.domains.converters.package_limits import (
    PackageTextLimits,
    PreparationLimit,
    bounded_join,
    check_package,
)
from airweave.domains.sync_pipeline.exceptions import SyncFailureError

if TYPE_CHECKING:
    from pptx.shapes.base import BaseShape
    from pptx.slide import Slide


def _extract_shape_text(shape: BaseShape) -> Iterator[str]:
    from pptx.shapes.graphfrm import GraphicFrame
    from pptx.shapes.group import GroupShape

    if isinstance(shape, GroupShape):
        for child in shape.shapes:
            yield from _extract_shape_text(child)
        return

    if shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            text = paragraph.text.strip()
            if text:
                yield text

    if shape.has_table:
        for row in cast(GraphicFrame, shape).table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                yield "| " + " | ".join(cells) + " |"


def _slide_content(slide: Slide) -> Iterator[str]:
    for shape in slide.shapes:
        yield from _extract_shape_text(shape)
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
        notes_text = slide.notes_slide.notes_text_frame.text.strip()
        if notes_text:
            yield f"> **Notes:** {notes_text}"


def _extract_slide(slide: Slide, slide_idx: int, limits: PackageTextLimits) -> str:
    content = _slide_content(slide)
    first = next(content, None)
    if first is None:
        return ""
    return bounded_join(chain((f"## Slide {slide_idx}", first), content), "\n\n", limits)


async def extract_pptx_text(path: str, limits: PackageTextLimits | None = None) -> Optional[str]:
    """Keep the text-only consumer contract used by oversized-PPTX OCR fallback."""
    return (await extract_pptx(path, limits)).text


async def extract_pptx(path: str, limits: PackageTextLimits | None = None) -> ConversionResult:
    """Extract local text with conservative disclosure of uninterpreted related content."""
    try:
        from pptx import Presentation
    except ImportError:
        raise SyncFailureError("python-pptx required for PPTX text extraction but not installed")

    limits = limits or PackageTextLimits()

    def _extract() -> ConversionResult:
        name = os.path.basename(path)

        try:
            check_package(path, limits)
            prs = Presentation(path)
        except PreparationLimit:
            raise
        except Exception as exc:
            logger.warning(f"Failed to open PPTX {name}: {exc}")
            return ConversionResult(text=None)

        slides = (
            _extract_slide(slide, index, limits) for index, slide in enumerate(prs.slides, start=1)
        )
        markdown = bounded_join((text for text in slides if text), "\n\n---\n\n", limits)
        gap = (
            "embedded_content_unprocessed"
            if any(_has_embedded_content(slide) for slide in prs.slides)
            else None
        )
        logger.debug(f"PPTX {name}: extracted {len(markdown)} chars")
        return ConversionResult(text=markdown or (None if gap else ""), gap=gap)

    return await asyncio.to_thread(_extract)


def _has_embedded_content(slide: Slide) -> bool:
    """Inspect used slide/story relationships, never follow external media links."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    unprocessed = {
        RT.IMAGE,
        RT.CHART,
        RT.DIAGRAM_DATA,
        RT.OLE_OBJECT,
        RT.AUDIO,
        RT.VIDEO,
        RT.MEDIA,
    }
    parts = {slide.part, slide.slide_layout.part, slide.slide_layout.slide_master.part}
    if slide.has_notes_slide:
        parts.add(slide.notes_slide.part)
        parts.add(slide.notes_slide.part.notes_master.part)
    return any(
        relationship.reltype in unprocessed for part in parts for relationship in part.rels.values()
    )
