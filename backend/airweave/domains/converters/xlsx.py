"""Bounded XLSX to markdown extraction using openpyxl's read-only reader."""

from collections.abc import Iterator
from pathlib import Path
from zipfile import ZipFile

from airweave.core.logging import logger
from airweave.domains.converters._base import BaseTextConverter, ConversionResult
from airweave.domains.converters.xlsx_limits import XlsxLimits
from airweave.domains.storage.limits import MAX_FILE_SIZE_BYTES
from airweave.domains.sync_pipeline.async_helpers import run_in_thread_pool
from airweave.domains.sync_pipeline.exceptions import EntityProcessingError, SyncFailureError


class XlsxPreparationLimit(EntityProcessingError):
    """The original exceeds the supported XLSX preparation content bounds."""


class XlsxConverter(BaseTextConverter):
    """Preserve formulas and multilingual text without trusting declared dimensions."""

    def __init__(self, limits: XlsxLimits | None = None) -> None:
        """Use one immutable limit policy for extraction and provenance."""
        self.limits = limits or XlsxLimits()

    async def convert_batch(self, file_paths: list[str]) -> dict[str, ConversionResult]:
        """Serialize local expansions within a batch; the shared executor owns threads."""
        try:
            import openpyxl  # noqa: F401
        except ImportError as exc:
            raise SyncFailureError("openpyxl package required for XLSX conversion") from exc

        results: dict[str, ConversionResult] = {}
        for path in file_paths:
            try:
                text = await run_in_thread_pool(self._extract, path)
                results[path] = ConversionResult(text=text)
            except XlsxPreparationLimit as exc:
                logger.warning(f"XLSX preparation limit: {exc}")
                results[path] = ConversionResult(text=None, failure_reason="preparation_limit")
            except Exception as exc:
                logger.warning(f"XLSX conversion failed: {exc}")
                results[path] = ConversionResult(text=None)
        return results

    def _check_package(self, path: str) -> None:
        """Bound eager shared-string/style parsing before openpyxl opens the package."""
        if Path(path).stat().st_size > MAX_FILE_SIZE_BYTES:
            raise XlsxPreparationLimit("original input bytes")
        with ZipFile(path) as package:
            members = package.infolist()
            if len(members) > self.limits.maximum_package_members:
                raise XlsxPreparationLimit("package member count")
            if any(member.file_size > self.limits.maximum_member_bytes for member in members):
                raise XlsxPreparationLimit("expanded package member bytes")
            if sum(member.file_size for member in members) > self.limits.maximum_expanded_bytes:
                raise XlsxPreparationLimit("expanded package bytes")

    def _extract(self, path: str) -> str:
        from openpyxl import load_workbook

        self._check_package(path)
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            if not workbook.sheetnames:
                raise EntityProcessingError("XLSX has no sheets")
            parts: list[str] = []
            output_bytes = 0
            total_rows = 0
            total_cells = 0

            def append(text: str) -> None:
                nonlocal output_bytes
                output_bytes += len(text.encode("utf-8")) + int(bool(parts))
                if output_bytes > self.limits.maximum_output_bytes:
                    raise XlsxPreparationLimit("extracted UTF-8 bytes")
                parts.append(text)

            for sheet in workbook.worksheets:
                # Ignore incorrect producer dimensions; actual sparse gaps still count below.
                sheet.reset_dimensions()
                rows: list[list[str]] = []
                width = 0
                for values in sheet.iter_rows(values_only=True):
                    total_rows += 1
                    if total_rows > self.limits.maximum_rows:
                        raise XlsxPreparationLimit("worksheet rows")
                    # openpyxl creates this tuple before yielding it. Explicit column names
                    # are bounded by its parser; malformed coordinate-less rows are only
                    # package-bounded before this check, not subject to a hard RSS limit.
                    width = max(width, len(values))
                    if width > self.limits.maximum_columns:
                        raise XlsxPreparationLimit("worksheet columns")
                    if total_cells + (len(rows) + 1) * width > self.limits.maximum_cells:
                        raise XlsxPreparationLimit("rendered worksheet cells")
                    rows.append(["" if value is None else str(value) for value in values])
                total_cells += len(rows) * width
                for line in self._sheet_lines(sheet.title, rows, width):
                    append(line)
            return "\n".join(parts)
        finally:
            workbook.close()

    @staticmethod
    def _sheet_lines(title: str, rows: list[list[str]], width: int) -> Iterator[str]:
        """Render the bounded final rectangle without truncating later wider rows."""
        yield f"## Sheet: {title}\n"
        if not rows:
            yield "*Empty sheet*\n"
        elif len(rows) == 1:
            for value in rows[0]:
                if value:
                    yield f"- {value}"
        else:
            for index, row in enumerate(rows):
                yield "| " + " | ".join(row + [""] * (width - len(row))) + " |"
                if index == 0:
                    yield "| " + " | ".join(["---"] * width) + " |"
        yield ""
