"""Tests for PptxConverter — _try_extract success and failure paths."""

from unittest.mock import AsyncMock, patch

import pytest

from airweave.domains.converters.pptx import PptxConverter


@pytest.fixture
def converter():
    return PptxConverter(ocr_provider=None)


class TestPptxConverter:
    @pytest.mark.asyncio
    async def test_try_extract_success(self, converter):
        """When extract_pptx_text returns content, _try_extract returns it."""
        with patch(
            "airweave.domains.converters.pptx.extract_pptx_text",
            new_callable=AsyncMock,
            return_value="# Slide 1\n\nBullet point",
        ):
            result = await converter._try_extract("/fake/pres.pptx")

        assert result == "# Slide 1\n\nBullet point"

    @pytest.mark.asyncio
    async def test_try_extract_returns_none(self, converter):
        """When extract_pptx_text returns None, _try_extract returns None (needs OCR)."""
        with patch(
            "airweave.domains.converters.pptx.extract_pptx_text",
            new_callable=AsyncMock,
            return_value=None,
        ):
            result = await converter._try_extract("/fake/pres.pptx")

        assert result is None


@pytest.mark.asyncio
@pytest.mark.parametrize("media", ["none", "picture", "chart", "group_picture", "layout"])
async def test_pptx_mixed_and_grouped_content_preserves_text_notes_original(tmp_path, media):
    from io import BytesIO

    from PIL import Image
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text = "नमस्ते — اردو — 中文. Retained PowerPoint text with enough useful context."
    slide.shapes.add_textbox(0, 0, Inches(5), Inches(1)).text = text
    group = slide.shapes.add_group_shape()
    nested = group.shapes.add_group_shape()
    nested.shapes.add_textbox(0, 0, Inches(5), Inches(1)).text = "Grouped 日本語 e\u0301"
    slide.notes_slide.notes_text_frame.text = "Retained notes: café 中文"
    image = BytesIO()
    Image.new("RGB", (10, 10), "blue").save(image, format="PNG")
    if media == "picture":
        slide.shapes.add_picture(image, 0, 0, Inches(1), Inches(1))
    elif media == "group_picture":
        nested.shapes.add_picture(image, 0, 0, Inches(1), Inches(1))
    elif media == "chart":
        data = CategoryChartData()
        data.categories = ["Chart only label"]
        data.add_series("Values", [42])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, Inches(2), Inches(2), data)
    elif media == "layout":
        # Conservatively disclose inherited linked media without fetching the URL.
        slide.slide_layout.part.relate_to(
            "https://never-fetch.invalid/background", RT.IMAGE, is_external=True
        )
    path = tmp_path / "mixed.pptx"
    presentation.save(path)
    before = path.read_bytes()
    ocr = AsyncMock()
    ocr.convert_batch.side_effect = AssertionError("Usable mixed text must not call OCR")
    result = (await PptxConverter(ocr).convert_batch([str(path)]))[str(path)]
    assert text in result.text and "Grouped 日本語 e\u0301" in result.text
    assert "Retained notes: café 中文" in result.text
    assert "Chart only label" not in result.text
    assert result.gap == (None if media == "none" else "embedded_content_unprocessed")
    ocr.convert_batch.assert_not_awaited()
    assert path.read_bytes() == before
