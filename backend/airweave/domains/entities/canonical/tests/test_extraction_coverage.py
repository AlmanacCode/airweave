"""Real SQL and conversion; synthetic originals, embeddings and remote feed."""

from unittest.mock import AsyncMock, MagicMock
from uuid import uuid4

import fitz
import pytest
from sqlalchemy import select, update

from airweave.adapters.storage.filesystem import FilesystemBackend
from airweave.core.logging import logger
from airweave.domains.converters.registry import ConverterRegistry
from airweave.domains.embedders.fakes.embedder import FakeDenseEmbedder, FakeSparseEmbedder
from airweave.domains.entities.canonical.extraction_models import (
    ExtractionCoverage,
    ExtractionOutcome,
)
from airweave.domains.entities.canonical.projection_models import ProjectionLocator
from airweave.domains.entities.canonical.projection_store import (
    CanonicalProjectionStore,
    current_extraction,
    publication_matches,
)
from airweave.domains.entities.canonical.projector import CanonicalProjector
from airweave.domains.entities.canonical.requests import RecordIdentity
from airweave.domains.entities.canonical.tests.helpers import bind_projection, capture, observation
from airweave.domains.entities.canonical.tests.test_gmail_projection import part
from airweave.domains.search.owned import OwnedSearchService
from airweave.domains.search.owned_models import OwnedSearchRequest
from airweave.domains.sync_pipeline.processors.chunk_embed import ChunkEmbedProcessor
from airweave.models import Entity, Sync
from airweave.models.projection_generation import ProjectionGeneration
from airweave.platform.destinations.vespa.transformer import EntityTransformer


def original(*attachments, completeness="complete"):
    return observation(
        identity=RecordIdentity(record_type="message", native_id="m1"),
        completeness=completeness,
        payload={
            "id": "m1",
            "threadId": "t1",
            "internalDate": "1000",
            "payload": {
                "mimeType": "multipart/mixed",
                "parts": [
                    part(b"Intact fundraising email body with useful context."),
                    *attachments,
                ],
            },
        },
    )


def destination(collection_id):
    result = MagicMock(collection_id=collection_id, feed_prepared=AsyncMock())
    result.prepare_documents = lambda chunks: {
        "base_entity": [
            EntityTransformer(collection_id=result.collection_id).transform(c) for c in chunks
        ]
    }
    return result


def projector(database, storage):
    return CanonicalProjector(
        CanonicalProjectionStore(),
        lambda _organization: database(),
        ChunkEmbedProcessor(ConverterRegistry(), FakeDenseEmbedder(), FakeSparseEmbedder()),
        storage,
    )


async def test_message_pdf_video_and_missing_part_publish_truthful_coverage(
    database, source, tmp_path
):
    service, fence = source
    binding = await bind_projection(database, fence, "gmail")
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text(
        (40, 40), "Retained PDF investment discussion and detailed fundraising plans. " * 4
    )
    data = pdf.tobytes()
    pdf.close()
    item = original(
        part(data, mime="application/pdf", filename="deck.pdf"),
        part(b"synthetic video bytes", mime="video/mp4", filename="meeting.mp4"),
        {
            "mimeType": "application/pdf",
            "filename": "large.pdf",
            "body": {"attachmentId": "missing", "size": 500000000},
        },
        completeness="partial",
    )
    item.payload["payload"]["parts"][0] = part(
        "Intact fundraising email body: नमस्ते — café".encode(),
        headers=[{"name": "Content-Type", "value": "text/plain; charset=gb2312"}],
    )
    await capture(database, service, fence, item)
    store = CanonicalProjectionStore()
    target = destination(binding.collection_id)
    result = await projector(database, FilesystemBackend(tmp_path)).batch(
        fence.organization_id, fence.sync_id, "gmail", target, logger
    )
    assert result.published == 1 and result.failed == 0
    async with database() as db:
        row = await db.scalar(select(Entity).where(Entity.sync_id == fence.sync_id))
        coverage = await current_extraction(
            db, fence.organization_id, fence.sync_id, row.id, row.record_revision
        )
        assert coverage.status == "partial"
        assert [p.outcome for p in coverage.parts] == [
            "indexed",
            "indexed",
            "unsupported",
            "unavailable_original",
        ]
        assert coverage.parts[1].key == "/payload/parts/1"
        recovery = coverage.parts[0].charset_recoveries[0]
        assert recovery.source_path == "/payload/parts/0"
        assert recovery.from_charset == "gb2312" and recovery.to_charset == "utf-8"
        assert (
            coverage.parts[1].kind == "file" and coverage.parts[1].media_type == "application/pdf"
        )
        assert row.completeness == "partial" and row.source_payload == item.payload
        assert not await store.pending(db, fence.organization_id, fence.sync_id)
        locator = ProjectionLocator(
            record_id=row.id,
            revision=row.record_revision,
            pipeline_version=row.indexed_pipeline_version,
            generation=row.indexed_generation,
            part_index=1,
        )
        assert (
            await db.scalar(select(Entity.id).join(Sync).where(publication_matches(locator)))
            == row.id
        )
        assert (
            await db.scalar(
                select(Entity.id)
                .join(Sync)
                .where(publication_matches(locator.model_copy(update={"part_index": 2})))
            )
            is None
        )
        ctx = MagicMock()
        ctx.organization.id = fence.organization_id
        counts = await OwnedSearchService._coverage(
            db, ctx, OwnedSearchRequest(query="fundraising", sync_ids=(fence.sync_id,))
        )
        assert counts[0].partially_indexed_records == 1 and counts[0].pending_records == 0
    # New native revision withdraws every old part and its coverage before new extraction.
    await capture(
        database, service, fence, original(part(b"new video", mime="video/mp4", filename="new.mp4"))
    )
    async with database() as db:
        assert (
            await db.scalar(select(Entity.id).join(Sync).where(publication_matches(locator)))
            is None
        )
        assert (
            await current_extraction(
                db, fence.organization_id, fence.sync_id, row.id, row.record_revision
            )
            is None
        )
    assert (
        await projector(database, FilesystemBackend(tmp_path)).batch(
            fence.organization_id, fence.sync_id, "gmail", target, logger
        )
    ).published == 1
    async with database() as db:
        assert (
            await db.scalar(select(Entity.id).join(Sync).where(publication_matches(locator)))
            is None
        )
        old = await db.get(ProjectionGeneration, locator.generation)
        assert old.retired_at is not None


async def test_supported_attachment_conversion_failure_and_feed_failure_stay_pending(
    database, source, tmp_path
):
    service, fence = source
    binding = await bind_projection(database, fence, "gmail")
    await capture(
        database,
        service,
        fence,
        original(part(b"broken binary\x00", mime="application/pdf", filename="broken.pdf")),
    )
    target = destination(binding.collection_id)
    project = projector(database, FilesystemBackend(tmp_path))
    assert (
        await project.batch(fence.organization_id, fence.sync_id, "gmail", target, logger)
    ).failed == 1
    target.feed_prepared.assert_awaited_once()
    async with database() as db:
        partial = await db.scalar(select(Entity).where(Entity.sync_id == fence.sync_id))
        coverage = await current_extraction(
            db, fence.organization_id, fence.sync_id, partial.id, partial.record_revision
        )
        assert [p.outcome for p in coverage.parts] == ["indexed", "failed"]
        assert coverage.parts[1].reason == "conversion_failed"
        assert partial.projection_error == "conversion_failed"
    await capture(
        database, service, fence, original(part(b"video", mime="video/mp4", filename="clip.mp4"))
    )
    target.feed_prepared.side_effect = ConnectionError("synthetic interruption")
    assert (
        await project.batch(fence.organization_id, fence.sync_id, "gmail", target, logger)
    ).failed == 1
    async with database() as db:
        row = await db.scalar(select(Entity).where(Entity.sync_id == fence.sync_id))
        assert row.indexed_revision != row.record_revision
        assert (
            await current_extraction(
                db, fence.organization_id, fence.sync_id, row.id, row.record_revision
            )
            is None
        )
        assert (
            len(await CanonicalProjectionStore().pending(db, fence.organization_id, fence.sync_id))
            == 1
        )
    target.feed_prepared.side_effect = None
    assert (
        await project.batch(fence.organization_id, fence.sync_id, "gmail", target, logger)
    ).published == 1


async def test_unsupported_only_publication_is_explicit_and_retries_on_pipeline_upgrade(
    database, source, tmp_path
):
    service, fence = source
    binding = await bind_projection(database, fence, "google_drive")
    from airweave.domains.entities.canonical.tests.test_gmail_projection import blob, record

    captured = record({}, sync_id=fence.sync_id)
    ref = blob(captured, b"video bytes")
    storage = FilesystemBackend(tmp_path)
    await storage.write_file(ref.key, b"video bytes")
    await capture(
        database,
        service,
        fence,
        observation(
            identity=RecordIdentity(record_type="file", native_id="video"),
            payload={"id": "video", "name": "movie.mp4", "mimeType": "video/mp4"},
            blobs=(ref,),
        ),
    )
    target = destination(binding.collection_id)
    assert (
        await projector(database, storage).batch(
            fence.organization_id, fence.sync_id, "google_drive", target, logger
        )
    ).published == 1
    async with database() as db:
        row = await db.scalar(select(Entity).where(Entity.sync_id == fence.sync_id))
        assert row.indexed_chunk_count > 0 and row.completeness == "complete"
        assert (
            await current_extraction(
                db, fence.organization_id, fence.sync_id, row.id, row.record_revision
            )
        ).status == "unavailable"
        await db.execute(
            update(Sync).where(Sync.id == fence.sync_id).values(index_pipeline_version=2)
        )
        await db.commit()
        assert (
            len(await CanonicalProjectionStore().pending(db, fence.organization_id, fence.sync_id))
            == 1
        )
        assert (
            await current_extraction(
                db, fence.organization_id, fence.sync_id, row.id, row.record_revision
            )
            is None
        )


def test_coverage_rejects_missing_duplicate_and_contradictory_parts():
    one = ExtractionOutcome(part_index=0, key="/body", kind="body", outcome="indexed")
    for parts in [(one, one), (one.model_copy(update={"part_index": 1}),)]:
        with pytest.raises(ValueError):
            ExtractionCoverage(parts=parts)
    with pytest.raises(ValueError):
        ExtractionOutcome(part_index=0, key="a", kind="file", outcome="unsupported", reason=None)


async def test_coverage_manifest_is_exact_immutable_and_capture_fences_publication(
    database, source
):
    from airweave.domains.entities.canonical.projection_models import (
        ProjectionDocument,
        scope_projection_document_id,
    )

    service, fence = source
    binding = await bind_projection(database, fence, "gmail")
    await capture(database, service, fence, original())
    store = CanonicalProjectionStore()
    async with database() as db:
        work = (await store.pending(db, fence.organization_id, fence.sync_id))[0]
    generation, collection = uuid4(), binding.collection_id
    locator = ProjectionLocator(
        record_id=work.record.id,
        revision=work.record.revision,
        pipeline_version=work.pipeline_version,
        generation=generation,
        part_index=0,
    )
    documents = (
        ProjectionDocument(
            schema_name="base_entity",
            document_id=scope_projection_document_id(
                fence.sync_id, collection, f"Entity_{locator.encode()}__chunk_0"
            ),
        ),
    )
    indexed = ExtractionOutcome(part_index=0, key="/body", kind="body", outcome="indexed")
    async with database() as db:
        with pytest.raises(ValueError, match="exactly"):
            await store.prepare(
                db, work, generation, collection, documents, coverage=ExtractionCoverage(parts=())
            )
        assert await store.prepare(
            db,
            work,
            generation,
            collection,
            documents,
            coverage=ExtractionCoverage(parts=(indexed,)),
        )
        with pytest.raises(ValueError, match="immutable"):
            await store.prepare(
                db,
                work,
                generation,
                collection,
                documents,
                coverage=ExtractionCoverage(
                    parts=(indexed.model_copy(update={"key": "different"}),)
                ),
            )
        assert (
            await current_extraction(
                db, fence.organization_id, fence.sync_id, work.record.id, work.record.revision
            )
            is None
        )
    await capture(
        database, service, fence, original(part(b"video", mime="video/mp4", filename="clip.mp4"))
    )
    async with database() as db:
        assert not await store.publish(db, work, generation, 1)
        assert (
            await current_extraction(
                db, fence.organization_id, fence.sync_id, work.record.id, work.record.revision
            )
            is None
        )


async def test_inline_image_without_ocr_and_configured_conversion_failure_are_distinct(
    database, source, tmp_path
):
    service, fence = source
    binding = await bind_projection(database, fence, "gmail")
    item = original(part(b"retained image bytes", mime="image/png", filename="inline.png"))
    item.payload["payload"]["mimeType"] = "multipart/related"
    await capture(database, service, fence, item)
    storage = FilesystemBackend(tmp_path)
    target = destination(binding.collection_id)
    result = await projector(database, storage).batch(
        fence.organization_id, fence.sync_id, "gmail", target, logger
    )
    assert result.published == 1 and result.failed == 0
    async with database() as db:
        row = await db.scalar(select(Entity).where(Entity.sync_id == fence.sync_id))
        coverage = await current_extraction(
            db, fence.organization_id, fence.sync_id, row.id, row.record_revision
        )
        assert coverage.status == "partial"
        assert [p.outcome for p in coverage.parts] == ["indexed", "unsupported"]
        assert coverage.parts[1].media_type == "image/png"
        assert row.indexed_chunk_count > 0
        assert row.source_payload == item.payload
        # Enabling OCR changes the pipeline; an actual converter failure must not
        # masquerade as an unsupported-format success or preserve stale coverage.
        await db.execute(
            update(Sync).where(Sync.id == fence.sync_id).values(index_pipeline_version=2)
        )
        await db.commit()
    ocr = MagicMock(convert_batch=AsyncMock(side_effect=lambda paths: dict.fromkeys(paths)))
    configured = CanonicalProjector(
        CanonicalProjectionStore(),
        lambda _organization: database(),
        ChunkEmbedProcessor(ConverterRegistry(ocr), FakeDenseEmbedder(), FakeSparseEmbedder()),
        storage,
    )
    target.feed_prepared.reset_mock()
    failed = await configured.batch(fence.organization_id, fence.sync_id, "gmail", target, logger)
    assert failed.failed == 1 and failed.published == 1
    ocr.convert_batch.assert_awaited_once()
    target.feed_prepared.assert_awaited_once()
    async with database() as db:
        coverage = await current_extraction(
            db, fence.organization_id, fence.sync_id, row.id, row.record_revision
        )
        assert [p.outcome for p in coverage.parts] == ["indexed", "failed"]
        assert coverage.parts[1].reason == "conversion_failed"


async def test_drive_metadata_only_reports_missing_original_then_new_bytes_are_indexed(
    database, source, tmp_path
):
    from airweave.domains.entities.canonical.tests.test_gmail_projection import blob, record

    service, fence = source
    binding = await bind_projection(database, fence, "google_drive")
    identity = RecordIdentity(record_type="file", native_id="retained-file")
    payload = {"id": identity.native_id, "name": "Notes.txt", "mimeType": "text/plain"}
    metadata = observation(identity=identity, payload=payload, completeness="metadata_only")
    await capture(database, service, fence, metadata)
    storage = FilesystemBackend(tmp_path)
    target = destination(binding.collection_id)
    project = projector(database, storage)
    assert (
        await project.batch(fence.organization_id, fence.sync_id, "google_drive", target, logger)
    ).published == 1
    async with database() as db:
        row = await db.scalar(select(Entity).where(Entity.sync_id == fence.sync_id))
        coverage = await current_extraction(
            db, fence.organization_id, fence.sync_id, row.id, row.record_revision
        )
        assert coverage.status == "unavailable" and row.indexed_chunk_count > 0
        assert coverage.parts[0].outcome == "unavailable_original"
        assert coverage.parts[0].reason == "original_not_captured"
    content = b"Recovered original text contains the complete available document."
    ref = blob(record({}, sync_id=fence.sync_id), content).model_copy(
        update={"media_type": "text/plain"}
    )
    await storage.write_file(ref.key, content)
    await capture(
        database, service, fence, observation(identity=identity, payload=payload, blobs=(ref,))
    )
    assert (
        await project.batch(fence.organization_id, fence.sync_id, "google_drive", target, logger)
    ).published == 1
    async with database() as db:
        row = await db.scalar(select(Entity).where(Entity.sync_id == fence.sync_id))
        assert row.indexed_chunk_count > 0
        previous = ProjectionLocator(
            record_id=row.id,
            revision=row.record_revision,
            pipeline_version=row.indexed_pipeline_version,
            generation=row.indexed_generation,
            part_index=0,
        )
    # Losing retained body availability withdraws the prior text generation.
    await capture(database, service, fence, metadata)
    assert (
        await project.batch(fence.organization_id, fence.sync_id, "google_drive", target, logger)
    ).published == 1
    async with database() as db:
        row = await db.scalar(select(Entity).where(Entity.sync_id == fence.sync_id))
        assert row.indexed_chunk_count > 0
        assert (
            await db.scalar(select(Entity.id).join(Sync).where(publication_matches(previous)))
            is None
        )
        assert (
            await current_extraction(
                db, fence.organization_id, fence.sync_id, row.id, row.record_revision
            )
        ).status == "unavailable"


async def test_pdf_partial_and_unavailable_ocr_are_published_without_losing_originals(
    database, source, tmp_path
):
    from airweave.domains.entities.canonical.tests.test_gmail_projection import blob, record

    service, fence = source
    binding = await bind_projection(database, fence, "google_drive")
    storage = FilesystemBackend(tmp_path)
    originals = {}
    for native_id, text in (
        ("mixed", "Preserved embedded text beside an unread scanned diagram."),
        ("scan", None),
        ("broken", None),
    ):
        pdf = fitz.open()
        page = pdf.new_page()
        if text:
            page.insert_text((40, 40), text)
        image = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 2, 2), False)
        image.clear_with(128)
        page.insert_image(fitz.Rect(40, 100, 80, 140), pixmap=image)
        content = pdf.tobytes() if native_id != "broken" else b"\x00\xff damaged PDF"
        pdf.close()
        ref = blob(record({}, sync_id=fence.sync_id), content).model_copy(
            update={"media_type": "application/pdf"}
        )
        await storage.write_file(ref.key, content)
        originals[native_id] = (ref, content)
        await capture(
            database,
            service,
            fence,
            observation(
                identity=RecordIdentity(record_type="file", native_id=native_id),
                payload={
                    "id": native_id,
                    "name": native_id + ".pdf",
                    "mimeType": "application/pdf",
                },
                blobs=(ref,),
            ),
        )
    result = await projector(database, storage).batch(
        fence.organization_id,
        fence.sync_id,
        "google_drive",
        destination(binding.collection_id),
        logger,
    )
    assert result.published == 3 and result.failed == 1
    async with database() as db:
        rows = (await db.scalars(select(Entity).where(Entity.sync_id == fence.sync_id))).all()
        by_native = {row.source_payload["id"]: row for row in rows}
        for native, status in (
            ("mixed", "partial"),
            ("scan", "unavailable"),
            ("broken", "unavailable"),
        ):
            row = by_native[native]
            coverage = await current_extraction(
                db, fence.organization_id, fence.sync_id, row.id, row.record_revision
            )
            assert coverage.status == status
            if native == "mixed":
                assert coverage.parts[0].gaps == ("ocr_unavailable",)
                assert row.indexed_chunk_count > 0
            else:
                assert coverage.parts[0].reason == (
                    "conversion_failed" if native == "broken" else "ocr_unavailable"
                )
                assert row.indexed_chunk_count > 0
                assert coverage.parts[-1].kind == "metadata"
            ref, original = originals[native]
            assert await storage.read_file(ref.key) == original


async def test_inert_attachment_decode_failure_publishes_body_with_failed_part(
    database, source, tmp_path
):
    service, fence = source
    binding = await bind_projection(database, fence, "gmail")
    await capture(
        database,
        service,
        fence,
        original(
            part(b"BEGIN:VCALENDAR\r\nSUMMARY:Review\r\nEND:VCALENDAR\r\n", mime="text/calendar"),
            part(b"\xff", mime="text/calendar"),
        ),
    )
    target = destination(binding.collection_id)
    result = await projector(database, FilesystemBackend(tmp_path)).batch(
        fence.organization_id, fence.sync_id, "gmail", target, logger
    )
    assert result.published == 1 and result.failed == 1
    target.feed_prepared.assert_awaited_once()
    async with database() as db:
        row = await db.scalar(select(Entity).where(Entity.sync_id == fence.sync_id))
        coverage = await current_extraction(
            db, fence.organization_id, fence.sync_id, row.id, row.record_revision
        )
        assert coverage.status == "partial"
        assert [p.outcome for p in coverage.parts] == ["indexed", "indexed", "failed"]
        assert coverage.parts[2].reason == "conversion_failed"
        assert row.projection_error == "conversion_failed"


async def test_xlsx_limit_keeps_published_mail_body_and_captured_original(
    database, source, tmp_path
):
    """Real XLSX extraction + SQL; synthetic MIME and fake embeddings, no provider calls."""
    from io import BytesIO

    from openpyxl import Workbook

    from airweave.domains.entities.canonical.mail_body import current_mail_body

    workbook = Workbook()
    workbook.active["A1"] = "नमस्ते"
    workbook.active["XFD1"] = "Original retained far column"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    data = stream.getvalue()
    service, fence = source
    binding = await bind_projection(database, fence, "gmail")
    item = original(
        part(
            data,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename="sparse.xlsx",
        )
    )
    await capture(database, service, fence, item)
    target = destination(binding.collection_id)
    result = await projector(database, FilesystemBackend(tmp_path)).batch(
        fence.organization_id, fence.sync_id, "gmail", target, logger
    )
    assert result.failed == 1 and result.published == 1
    target.feed_prepared.assert_awaited_once()
    async with database() as db:
        row = (await db.scalars(select(Entity).where(Entity.sync_id == fence.sync_id))).one()
        coverage = await current_extraction(
            db, fence.organization_id, fence.sync_id, row.id, row.record_revision
        )
        assert coverage.status == "partial"
        assert [p.outcome for p in coverage.parts] == ["indexed", "failed"]
        assert coverage.parts[1].reason == "preparation_limit"
        captured = await service.store.read(db, fence.organization_id, fence.sync_id, row.id)
        assert captured.completeness == "complete" and captured.content_access == "available"
        assert captured.payload == item.payload
        body = await db.scalar(
            select(current_mail_body().with_only_columns(ProjectionGeneration.mail_body_text)
                   .scalar_subquery()).select_from(Entity).join(Sync, Sync.id == Entity.sync_id)
        )
        assert "intact fundraising email body with useful context" in body.lower()


@pytest.mark.parametrize("over_limit", [False, True])
async def test_embedded_content_gap_reaches_partial_mail_coverage_offline(tmp_path, over_limit):
    """Actual MIME mapping/converters/coverage; no SQL, embedding, OCR or remote feed."""
    from io import BytesIO
    from types import SimpleNamespace

    from docx import Document
    from PIL import Image

    from airweave.domains.converters.docx import DocxConverter
    from airweave.domains.converters.package_limits import PackageTextLimits
    from airweave.domains.entities.canonical.gmail_projection import map_gmail
    from airweave.domains.entities.canonical.projection_models import (
        ProjectionBinding,
        ProjectionWork,
    )
    from airweave.domains.entities.canonical.projector import (
        ProjectionConversionTracker,
        _conversion_coverage,
        _select_inputs,
    )
    from airweave.domains.entities.canonical.tests.test_gmail_projection import record
    from airweave.domains.sync_pipeline.pipeline.text_builder import TextualRepresentationBuilder

    document = Document()
    text = "नमस्ते — اردو — 中文. Retained attachment paragraph with useful source context."
    document.add_paragraph(text)
    image = BytesIO()
    Image.new("RGB", (10, 10), "blue").save(image, format="PNG")
    document.add_picture(image)
    content = BytesIO()
    document.save(content)
    item = record(
        {
            "mimeType": "multipart/mixed",
            "parts": [
                part(b"Intact readable parent email body with enough useful source context."),
                part(
                    content.getvalue(),
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    filename="mixed.docx",
                ),
            ],
        }
    )
    before = item.model_dump()
    mapped = await map_gmail(item, AsyncMock(), tmp_path)
    work = ProjectionWork(
        binding=ProjectionBinding(
            source_connection_id=uuid4(), source_name="gmail", collection_id=uuid4()
        ),
        organization_id=uuid4(),
        record=item,
        pipeline_version=1,
        previous_generation=None,
    )
    selected, coverage = _select_inputs(mapped, work, "gmail", uuid4(), lambda _: True)
    registry = ConverterRegistry()
    if over_limit:
        # Narrow fixture dependency: real DOCX extraction with a small output budget.
        bounded_docx = DocxConverter(limits=PackageTextLimits(maximum_output_bytes=1))
        configured = MagicMock()
        configured.for_extension.side_effect = lambda extension, actual=registry: (
            bounded_docx if extension == ".docx" else actual.for_extension(extension)
        )
        configured.for_web.side_effect = registry.for_web
        registry = configured
    batch = await TextualRepresentationBuilder(registry).build_with_text(
        selected,
        SimpleNamespace(source_short_name="gmail", logger=MagicMock()),
        SimpleNamespace(entity_tracker=ProjectionConversionTracker(allow_failures=True)),
        strict_conversion=True,
        native_bodies={
            p.entity.entity_id: p.native_body
            for p in mapped.parts
            if p.entity is not None and p.native_body is not None
        },
    )
    covered = _conversion_coverage(batch, coverage)
    assert covered.status == "partial"
    assert [p.outcome for p in covered.parts] == ["indexed", "failed" if over_limit else "indexed"]
    if over_limit:
        assert covered.parts[1].reason == "preparation_limit"
        assert len(batch.representations) == 1
    else:
        assert covered.parts[1].gaps == ("embedded_content_unprocessed",)
        assert text in batch.representations[1].text
    assert "readable parent email body" in batch.representations[0].text.lower()
    assert item.model_dump() == before
